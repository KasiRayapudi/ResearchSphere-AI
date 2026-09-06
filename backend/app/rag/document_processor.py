"""
Document Processing Service - Real text extraction from PDF, DOCX, TXT, MD
"""
import os
import logging
from pathlib import Path
from typing import List, Tuple, Optional

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> Tuple[str, int]:
    """Extract text from PDF using PyMuPDF (fitz). Returns (text, page_count)."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(file_path)
        pages_text = []
        for page_num, page in enumerate(doc):
            text = page.get_text("text")
            if text.strip():
                pages_text.append(f"[Page {page_num + 1}]\n{text}")
        doc.close()
        full_text = "\n\n".join(pages_text)
        return full_text, len(doc)
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        raise ValueError(f"Failed to extract text from PDF: {e}")


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        # Also extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs)
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        raise ValueError(f"Failed to extract text from DOCX: {e}")


def extract_text_from_txt(file_path: str) -> str:
    """Extract text from TXT or Markdown files."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    except Exception as e:
        raise ValueError(f"Failed to read text file: {e}")


def extract_text(file_path: str, file_type: str) -> str:
    """Route extraction based on file type."""
    file_type = file_type.lower().strip(".")
    
    if file_type == "pdf":
        text, _ = extract_text_from_pdf(file_path)
        return text
    elif file_type == "docx":
        return extract_text_from_docx(file_path)
    elif file_type in ("txt", "md", "markdown", "csv"):
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def clean_text(text: str) -> str:
    """Clean extracted text: normalize whitespace, remove junk."""
    import re
    # Normalize whitespace
    text = re.sub(r'\r\n', '\n', text)
    text = re.sub(r'\r', '\n', text)
    # Remove excessive blank lines
    text = re.sub(r'\n{3,}', '\n\n', text)
    # Remove null bytes
    text = text.replace('\x00', '')
    return text.strip()


def chunk_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> List[dict]:
    """
    Split text into overlapping chunks using LangChain RecursiveCharacterTextSplitter.
    Returns list of dicts with 'content' and 'index'.
    """
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
        length_function=len,
    )
    
    chunks = splitter.split_text(text)
    return [
        {
            "content": chunk.strip(),
            "index": idx,
            "token_count": len(chunk.split()),
        }
        for idx, chunk in enumerate(chunks)
        if chunk.strip()
    ]
